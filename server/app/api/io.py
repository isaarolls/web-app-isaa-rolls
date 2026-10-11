"""Importazione (con verifica preventiva) ed esportazione dei dati."""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
import unicodedata
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from starlette.responses import Response

from .. import db
from ..common import D, q2, today
from ..http import Ctx, api, audit, bad
from ..services import exporting, stock
from ..services import expenses as exp_svc
from ..services import sales as sales_svc

# ---------------------------------------------------------------- lettura file


def _norm(h) -> str:
    h = unicodedata.normalize("NFKD", str(h or "")).encode("ascii", "ignore").decode().lower().strip()
    return re.sub(r"[^a-z0-9]+", "_", h).strip("_")


SYN = {
    "sku": ["sku", "codice", "codice_articolo", "cod", "code"],
    "name": ["nome", "name", "articolo", "prodotto", "descrizione_articolo"],
    "category": ["categoria", "category"],
    "kind": ["tipo", "kind", "tipologia"],
    "unit": ["unita", "unita_di_misura", "um", "u_m", "unit"],
    "min_qty": ["scorta_minima", "minimo", "min_qty", "quantita_minima"],
    "vat_rate": ["iva", "aliquota", "aliquota_iva", "iva_perc", "vat"],
    "location": ["posizione", "ubicazione", "location"],
    "qty": ["giacenza", "giacenza_iniziale", "quantita", "qty", "qta"],
    "unit_cost": ["costo", "costo_unitario", "prezzo_acquisto", "unit_cost"],
    "lot": ["lotto", "lot"],
    "expiry": ["scadenza", "data_scadenza", "expiry"],
    "supplier": ["fornitore", "supplier"],
    "price": ["prezzo", "prezzo_vendita", "prezzo_al_pubblico", "price"],
    "date": ["data", "data_documento", "date", "giorno"],
    "description": ["descrizione", "description", "causale"],
    "doc_number": ["numero", "numero_documento", "n_documento", "doc_number", "fattura"],
    "net": ["imponibile", "netto", "net"],
    "vat_amount": ["importo_iva", "iva_importo", "vat_amount"],
    "total": ["totale", "importo", "importo_totale", "lordo", "total", "importo_lordo"],
    "due_date": ["scadenza_pagamento", "data_scadenza_pagamento", "due_date"],
    "paid": ["pagata", "pagato", "paid"],
    "paid_at": ["data_pagamento", "pagata_il", "paid_at"],
    "channel": ["canale", "channel"],
    "method": ["metodo", "metodo_pagamento", "pagamento", "method"],
    "product": ["prodotto", "product", "codice_prodotto"],
    "pending": ["da_incassare", "in_attesa", "pending"],
}


def read_table(f: dict) -> list[dict]:
    name = (f["filename"] or "").lower()
    data = f["data"]
    rows: list[list] = []
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    elif name.endswith((".csv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4000], delimiters=";,\t")
            delim = dialect.delimiter
        except csv.Error:
            delim = ";"
        rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    else:
        raise bad("Formato non supportato: usa un file .csv o .xlsx")
    rows = [r for r in rows if any(str(x).strip() for x in r if x is not None)]
    if len(rows) < 2:
        raise bad("Il file non contiene righe di dati")
    headers = [_norm(h) for h in rows[0]]
    mapping = {}
    for i, h in enumerate(headers):
        for key, syns in SYN.items():
            if h in syns and key not in mapping.values():
                mapping[i] = key
                break
    out = []
    for n, r in enumerate(rows[1:], start=2):
        rec = {"_row": n}
        for i, key in mapping.items():
            v = r[i] if i < len(r) else None
            rec[key] = v.strip() if isinstance(v, str) else v
        out.append(rec)
    if len(out) > 20000:
        raise bad("File troppo grande (massimo 20.000 righe)")
    return out


def _num(v, field, errors, required=False, allow_neg=False):
    if v in (None, ""):
        if required:
            errors.append(f"{field} mancante")
        return None
    try:
        x = D(v)
    except Exception:
        errors.append(f"{field} non è un numero valido")
        return None
    if x < 0 and not allow_neg:
        errors.append(f"{field} negativo")
    return x


def _date(v, field, errors, required=False):
    if v in (None, ""):
        if required:
            errors.append(f"{field} mancante")
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(s[:10] if fmt == "%Y-%m-%d" else s, fmt).date()
        except ValueError:
            pass
    errors.append(f"{field} non è una data valida (usa gg/mm/aaaa)")
    return None


def _bool(v):
    return str(v or "").strip().lower() in ("si", "sì", "s", "yes", "y", "1", "true", "x", "vero")


KINDS = {"materia prima": "materia_prima", "materia_prima": "materia_prima", "semilavorato": "semilavorato",
         "confezione": "confezione", "packaging": "confezione", "rivendita": "rivendita", "consumabile": "consumabile"}


def validate(c: db.Cursor, kind: str, recs: list[dict]) -> list[dict]:
    units = {r["code"].lower(): r["code"] for r in c.all("SELECT code FROM units")}
    seen = set()
    out = []
    if kind == "items":
        for r in recs:
            e, w = [], []
            sku, name = str(r.get("sku") or "").strip(), str(r.get("name") or "").strip()
            if not sku:
                e.append("codice mancante")
            if not name:
                e.append("nome mancante")
            unit = units.get(str(r.get("unit") or "").strip().lower())
            if not unit:
                e.append(f"unità di misura «{r.get('unit') or ''}» non riconosciuta")
            k = KINDS.get(str(r.get("kind") or "materia prima").strip().lower())
            if not k:
                e.append("tipo non riconosciuto (materia prima, semilavorato, confezione, rivendita, consumabile)")
            qty = _num(r.get("qty"), "giacenza", e)
            cost = _num(r.get("unit_cost"), "costo", e)
            if qty and qty > 0 and cost is None:
                w.append("giacenza senza costo: sarà valorizzata a 0 €")
            dup = c.val("SELECT id FROM items WHERE lower(sku)=lower(%s)", [sku]) if sku else None
            if sku.lower() in seen:
                e.append("codice ripetuto nel file")
            seen.add(sku.lower())
            out.append({"row": r["_row"], "data": {
                "sku": sku, "name": name, "category": r.get("category"), "kind": k, "unit_code": unit,
                "min_qty": _num(r.get("min_qty"), "scorta minima", e) or 0, "vat_rate": _num(r.get("vat_rate"), "IVA", e),
                "location": r.get("location"), "qty": qty, "unit_cost": cost, "lot": r.get("lot"),
                "expiry": _date(r.get("expiry"), "scadenza", e), "supplier": r.get("supplier")},
                "errors": e, "warnings": w, "duplicate": bool(dup)})
    elif kind == "products":
        for r in recs:
            e, w = [], []
            name = str(r.get("name") or r.get("product") or "").strip()
            if not name:
                e.append("nome mancante")
            price = _num(r.get("price") or r.get("total"), "prezzo", e, required=True)
            code = str(r.get("sku") or "").strip() or None
            dup = c.val("SELECT id FROM products WHERE (%s::text IS NOT NULL AND lower(code)=lower(%s)) OR lower(name)=lower(%s)",
                        [code, code, name])
            if name.lower() in seen:
                e.append("prodotto ripetuto nel file")
            seen.add(name.lower())
            out.append({"row": r["_row"], "data": {"code": code, "name": name, "category": r.get("category"), "price": price,
                                                   "vat_rate": _num(r.get("vat_rate"), "IVA", e)},
                        "errors": e, "warnings": w, "duplicate": bool(dup)})
    elif kind == "expenses":
        cats = {x["name"].lower(): x["id"] for x in c.all("SELECT id, name FROM expense_categories")}
        for r in recs:
            e, w = [], []
            d = _date(r.get("date"), "data", e, required=True)
            desc = str(r.get("description") or "").strip()
            if not desc:
                e.append("descrizione mancante")
            cat = cats.get(str(r.get("category") or "").strip().lower())
            if not cat:
                e.append(f"categoria «{r.get('category') or ''}» inesistente")
            net = _num(r.get("net"), "imponibile", e, allow_neg=True)
            vat = _num(r.get("vat_amount"), "importo IVA", e, allow_neg=True)
            tot = _num(r.get("total"), "totale", e, allow_neg=True)
            rate = _num(r.get("vat_rate"), "aliquota", e)
            if net is None and tot is None:
                e.append("indica imponibile o totale")
            if net is not None and vat is not None and tot is not None and q2(net + vat) != q2(tot):
                e.append("imponibile + IVA diverso dal totale")
            sup = None
            if r.get("supplier"):
                sup = c.val("SELECT id FROM suppliers WHERE lower(legal_name)=lower(%s) OR lower(trade_name)=lower(%s)",
                            [r["supplier"], r["supplier"]])
                if not sup:
                    w.append(f"fornitore «{r['supplier']}» verrà creato")
            paid = _bool(r.get("paid"))
            data = {"date": d, "description": desc, "category_id": cat, "supplier": r.get("supplier"), "supplier_id": sup,
                    "doc_number": r.get("doc_number"), "net": net, "vat": vat, "total": tot, "rate": rate,
                    "due_date": _date(r.get("due_date"), "scadenza", e), "paid": paid,
                    "paid_at": _date(r.get("paid_at"), "data pagamento", e) or d}
            key = (d, desc.lower(), str(tot or net), str(r.get("doc_number") or ""))
            if key in seen:
                e.append("riga duplicata nel file")
            seen.add(key)
            dup = None
            if d and not e:
                amount = tot if tot is not None else net
                dup = c.val("""SELECT id FROM expenses WHERE voided_at IS NULL AND doc_date=%s
                               AND (total_amount=%s OR net_amount=%s)
                               AND (lower(description)=lower(%s) OR (%s::text IS NOT NULL AND doc_number=%s))""",
                            [d, amount, amount, desc, data["doc_number"], data["doc_number"]])
            out.append({"row": r["_row"], "data": data, "errors": e, "warnings": w, "duplicate": bool(dup)})
    elif kind == "sales":
        chans = {x["name"].lower(): x["id"] for x in c.all("SELECT id, name FROM sales_channels")}
        methods = {x["name"].lower(): x for x in c.all("SELECT * FROM payment_methods")}
        for r in recs:
            e, w = [], []
            d = _date(r.get("date"), "data", e, required=True)
            ch = chans.get(str(r.get("channel") or "vendita diretta").strip().lower())
            if not ch:
                e.append(f"canale «{r.get('channel')}» inesistente")
            m = None
            if r.get("method"):
                m = methods.get(str(r["method"]).strip().lower())
                if not m:
                    e.append(f"metodo di pagamento «{r['method']}» inesistente")
            prod = None
            pref = r.get("product")
            if pref:
                prod = c.one("SELECT id, sale_price FROM products WHERE lower(code)=lower(%s) OR lower(name)=lower(%s)",
                             [str(pref), str(pref)])
                if not prod:
                    e.append(f"prodotto «{pref}» inesistente")
            qty = _num(r.get("qty"), "quantità", e) or Decimal(1)
            tot = _num(r.get("total") or r.get("price"), "importo", e, required=not prod)
            dup = None
            if d and ch and tot is not None and not pref:
                dup = c.val("SELECT id FROM sales WHERE status='valida' AND business_date=%s AND channel_id=%s AND gross_total=%s",
                            [d, ch, tot])
            out.append({"row": r["_row"], "data": {"date": d, "channel_id": ch, "method_id": m["id"] if m else None,
                                                   "product_id": prod["id"] if prod else None, "qty": qty, "total": tot,
                                                   "vat_rate": _num(r.get("vat_rate"), "IVA", e),
                                                   "pending": _bool(r.get("pending")), "description": r.get("description")},
                        "errors": e, "warnings": w, "duplicate": bool(dup)})
    else:
        raise bad("Tipo di importazione non valido")
    return out


IMPORT_PERMS = {"items": "inventory.edit", "products": "products.edit", "expenses": "expenses.edit", "sales": "sales.create"}


def _get_rows(ctx):
    kind = ctx.params["kind"]
    if kind not in IMPORT_PERMS:
        raise bad("Tipo di importazione non valido")
    ctx.require("import.run")
    ctx.require(IMPORT_PERMS[kind])
    f = ctx.files.get("file")
    if not f:
        raise bad("Nessun file ricevuto")
    return kind, read_table(f)


def _summary(rows):
    return {"total": len(rows), "errors": sum(1 for r in rows if r["errors"]),
            "duplicates": sum(1 for r in rows if r["duplicate"] and not r["errors"]),
            "valid": sum(1 for r in rows if not r["errors"] and not r["duplicate"])}


@api("POST", "/import/{kind}/preview")
def import_preview(ctx: Ctx):
    kind, recs = _get_rows(ctx)
    with db.read() as c:
        rows = validate(c, kind, recs)
    return {"summary": _summary(rows), "rows": rows[:2000]}


@api("POST", "/import/{kind}/commit")
def import_commit(ctx: Ctx):
    kind, recs = _get_rows(ctx)
    skip_errors = ctx.form.get("skip_errors") == "1"
    unload = ctx.form.get("unload_stock") == "1"
    with db.tx() as c:
        rows = validate(c, kind, recs)
        summ = _summary(rows)
        if summ["errors"] and not skip_errors:
            raise bad(f"Il file contiene {summ['errors']} righe con errori: correggile o scegli di importare solo le righe valide")
        todo = [r for r in rows if not r["errors"] and not r["duplicate"]]
        n = 0
        for r in todo:
            d = r["data"]
            if kind == "items":
                cat = None
                if d["category"]:
                    cat = c.val("SELECT id FROM item_categories WHERE lower(name)=lower(%s)", [d["category"]]) or \
                        c.val("INSERT INTO item_categories(name) VALUES (%s) RETURNING id", [d["category"]])
                sup = c.val("SELECT id FROM suppliers WHERE lower(legal_name)=lower(%s)", [d["supplier"]]) if d["supplier"] else None
                iid = c.val("""INSERT INTO items(sku,name,kind,category_id,unit_code,min_qty,vat_rate,location,default_supplier_id,track_lots)
                               VALUES (%s,%s,%s,%s,%s,%s,COALESCE(%s,10),%s,%s,%s) RETURNING id""",
                            [d["sku"], d["name"], d["kind"], cat, d["unit_code"], d["min_qty"], d["vat_rate"], d["location"], sup,
                             bool(d["lot"] or d["expiry"])])
                if d["qty"] and d["qty"] > 0:
                    stock.move(c, ctx.uid, iid, d["qty"], "giacenza_iniziale", unit_cost=d["unit_cost"] or 0,
                               reason="Importazione giacenze iniziali", lot_code=d["lot"], expiry_date=d["expiry"])
            elif kind == "products":
                cat = None
                if d["category"]:
                    cat = c.val("SELECT id FROM product_categories WHERE lower(name)=lower(%s)", [d["category"]]) or \
                        c.val("INSERT INTO product_categories(name) VALUES (%s) RETURNING id", [d["category"]])
                c.exec("INSERT INTO products(code,name,category_id,sale_price,vat_rate) VALUES (%s,%s,%s,%s,COALESCE(%s,10))",
                       [d["code"], d["name"], cat, d["price"], d["vat_rate"]])
            elif kind == "expenses":
                sup = d["supplier_id"]
                if not sup and d["supplier"]:
                    sup = c.val("INSERT INTO suppliers(legal_name) VALUES (%s) RETURNING id", [d["supplier"]])
                e = exp_svc.create(c, ctx.uid, {
                    "description": d["description"], "category_id": d["category_id"], "supplier_id": sup,
                    "doc_number": d["doc_number"], "doc_date": d["date"], "net_amount": d["net"], "vat_amount": d["vat"],
                    "total_amount": d["total"], "vat_rate": d["rate"], "due_date": d["due_date"],
                    "doc_type": "fattura" if d["doc_number"] else "altro",
                })
                if d["paid"]:
                    exp_svc.add_payment(c, ctx.uid, e["id"], {"paid_at": d["paid_at"], "amount": None})
            elif kind == "sales":
                when = dt.datetime.combine(d["date"], dt.time(12, 0)).isoformat()
                if d["product_id"]:
                    line = {"product_id": d["product_id"], "qty": d["qty"]}
                    if d["total"] is not None:
                        line["unit_price"] = q2(d["total"] / d["qty"])
                    data = {"sold_at": when, "channel_id": d["channel_id"], "mode": "analitica", "lines": [line],
                            "skip_stock": not unload}
                else:
                    data = {"sold_at": when, "channel_id": d["channel_id"], "mode": "riepilogo",
                            "lines": [{"description": d["description"] or "Incasso importato", "qty": 1,
                                       "unit_price": d["total"], "vat_rate": d["vat_rate"]}]}
                s = sales_svc.create_sale(c, ctx.uid, ctx.perms | {"sales.void"}, data)
                if d["method_id"]:
                    c.exec("DELETE FROM sale_payments WHERE sale_id=%s", [s["id"]])
                    sales_svc._add_payments(c, s, [{"method_id": d["method_id"], "amount": s["gross_total"],
                                                    "pending": True if d["pending"] else None}], None, s["gross_total"])
            n += 1
        audit(c, ctx, "import", kind, None, f"Importazione {kind}: {n} righe importate, {summ['duplicates']} duplicati saltati, "
                                            f"{summ['errors']} righe con errori escluse")
    return {"imported": n, "skipped_duplicates": summ["duplicates"], "skipped_errors": summ["errors"] if skip_errors else 0}


@api("GET", "/import/{kind}/template")
def import_template(ctx: Ctx):
    kind = ctx.params["kind"]
    templates = {
        "items": (["codice", "nome", "categoria", "tipo", "unita", "scorta_minima", "iva", "posizione", "giacenza", "costo_unitario", "lotto", "scadenza", "fornitore"],
                  [["FAR-00", "Farina 00", "Farine e impasti", "materia prima", "kg", "10", "4", "Scaffale A", "25", "0,85", "", "", ""]]),
        "products": (["codice", "nome", "categoria", "prezzo", "iva"], [["ROLL-CLA", "Roll classico", "Rolls", "4,50", "10"]]),
        "expenses": (["data", "descrizione", "categoria", "fornitore", "numero", "imponibile", "importo_iva", "totale", "scadenza_pagamento", "pagata", "data_pagamento"],
                     [["01/09/2026", "Affitto settembre", "Affitto", "Immobiliare Rossi", "12", "1000,00", "220,00", "1220,00", "05/09/2026", "si", "05/09/2026"]]),
        "sales": (["data", "canale", "metodo", "prodotto", "quantita", "importo", "iva", "descrizione", "da_incassare"],
                  [["15/09/2026", "Vendita diretta", "Contanti", "", "", "350,00", "10", "Incassi del giorno", ""],
                   ["15/09/2026", "Glovo", "Pagato su piattaforma (da accreditare)", "Roll classico", "12", "", "", "", "si"]]),
    }
    if kind not in templates:
        raise bad("Modello non disponibile")
    cols, ex = templates[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = "Dati"
    ws.append(cols)
    for r in ex:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="modello-{kind}.xlsx"'})


# ---------------------------------------------------------------- esportazioni

EXPORTS = {
    "items": ("inventory.view", "Articoli di magazzino", """
        SELECT i.sku, i.name, i.kind, ic.name AS category, i.unit_code, i.qty_on_hand, i.min_qty, i.avg_cost,
               GREATEST(i.qty_on_hand,0)*i.avg_cost AS stock_value, i.last_purchase_price, i.last_purchase_date, i.location,
               s.legal_name AS supplier, i.is_active
        FROM items i LEFT JOIN item_categories ic ON ic.id=i.category_id LEFT JOIN suppliers s ON s.id=i.default_supplier_id ORDER BY i.name""",
        [("sku", "Codice"), ("name", "Nome"), ("kind", "Tipo"), ("category", "Categoria"), ("unit_code", "U.M."),
         ("qty_on_hand", "Giacenza"), ("min_qty", "Scorta minima"), ("avg_cost", "Costo medio"), ("stock_value", "Valore"),
         ("last_purchase_price", "Ultimo prezzo"), ("last_purchase_date", "Ultimo acquisto"), ("location", "Posizione"),
         ("supplier", "Fornitore"), ("is_active", "Attivo")], False),
    "movements": ("inventory.view", "Movimenti di magazzino", """
        SELECT m.id, m.at, i.sku, i.name, m.type, m.qty, i.unit_code, m.unit_cost, m.value, m.qty_after, m.reason,
               m.waste_reason, m.ref_type, m.ref_id, m.reversal_of, u.full_name AS user_name
        FROM stock_movements m JOIN items i ON i.id=m.item_id LEFT JOIN users u ON u.id=m.user_id
        WHERE (m.at AT TIME ZONE 'Europe/Rome')::date BETWEEN %s AND %s ORDER BY m.at, m.id""",
        [("id", "ID"), ("at", "Data e ora"), ("sku", "Codice"), ("name", "Articolo"), ("type", "Tipo"), ("qty", "Quantità"),
         ("unit_code", "U.M."), ("unit_cost", "Costo unitario"), ("value", "Valore"), ("qty_after", "Giacenza dopo"),
         ("reason", "Causale"), ("waste_reason", "Motivo spreco"), ("ref_type", "Documento"), ("ref_id", "ID documento"),
         ("reversal_of", "Storno di"), ("user_name", "Utente")], True),
    "sales": ("sales.view", "Vendite", """
        SELECT s.number, s.sold_at, s.business_date, s.kind, s.mode, ch.name AS channel, s.gross_before_discount, s.discount_total,
               s.gross_total, s.net_total, s.vat_total, s.commission_gross, s.status, s.void_reason, s.notes, u.full_name AS user_name
        FROM sales s JOIN sales_channels ch ON ch.id=s.channel_id LEFT JOIN users u ON u.id=s.user_id
        WHERE s.business_date BETWEEN %s AND %s ORDER BY s.sold_at""",
        [("number", "Numero"), ("sold_at", "Data e ora"), ("business_date", "Giorno"), ("kind", "Tipo"), ("mode", "Modalità"),
         ("channel", "Canale"), ("gross_before_discount", "Lordo"), ("discount_total", "Sconti"), ("gross_total", "Totale"),
         ("net_total", "Imponibile"), ("vat_total", "IVA"), ("commission_gross", "Commissione"), ("status", "Stato"),
         ("void_reason", "Motivo annullamento"), ("notes", "Note"), ("user_name", "Utente")], True),
    "sale-lines": ("sales.view", "Righe di vendita", """
        SELECT s.number, s.business_date, ch.name AS channel, l.description, l.qty, l.unit_price, l.discount, l.line_total,
               l.vat_rate, l.line_net, l.line_vat, s.status
        FROM sale_lines l JOIN sales s ON s.id=l.sale_id JOIN sales_channels ch ON ch.id=s.channel_id
        WHERE s.business_date BETWEEN %s AND %s ORDER BY s.sold_at, l.id""",
        [("number", "Vendita"), ("business_date", "Giorno"), ("channel", "Canale"), ("description", "Prodotto"),
         ("qty", "Quantità"), ("unit_price", "Prezzo"), ("discount", "Sconto"), ("line_total", "Totale"), ("vat_rate", "IVA %"),
         ("line_net", "Imponibile"), ("line_vat", "IVA"), ("status", "Stato")], True),
    "expenses": ("expenses.view", "Spese", """
        SELECT e.id, e.doc_date, e.competence_date, e.description, ec.name AS category, s.legal_name AS supplier, e.doc_type,
               e.doc_number, e.net_amount, e.vat_amount, e.total_amount, e.vat_deductible, e.due_date, e.paid_amount, e.status,
               e.voided_at, e.notes
        FROM expenses e JOIN expense_categories ec ON ec.id=e.category_id LEFT JOIN suppliers s ON s.id=e.supplier_id
        WHERE e.doc_date BETWEEN %s AND %s ORDER BY e.doc_date""",
        [("id", "ID"), ("doc_date", "Data documento"), ("competence_date", "Competenza"), ("description", "Descrizione"),
         ("category", "Categoria"), ("supplier", "Fornitore"), ("doc_type", "Tipo doc."), ("doc_number", "Numero"),
         ("net_amount", "Imponibile"), ("vat_amount", "IVA"), ("total_amount", "Totale"), ("vat_deductible", "IVA detraibile"),
         ("due_date", "Scadenza"), ("paid_amount", "Pagato"), ("status", "Stato"), ("voided_at", "Annullata il"), ("notes", "Note")], True),
    "suppliers": ("suppliers.view", "Fornitori", "SELECT * FROM suppliers ORDER BY legal_name",
                  [("legal_name", "Ragione sociale"), ("trade_name", "Nome commerciale"), ("vat_number", "Partita IVA"),
                   ("email", "E-mail"), ("phone", "Telefono"), ("contact_person", "Referente"), ("address", "Indirizzo"),
                   ("city", "Città"), ("payment_terms", "Condizioni pagamento"), ("lead_time_days", "Giorni consegna"),
                   ("notes", "Note")], False),
}


@api("GET", "/export/{what}.{fmt}")
def export_table(ctx: Ctx):
    what, fmt = ctx.params["what"], ctx.params["fmt"]
    if what not in EXPORTS or fmt not in ("csv", "xlsx"):
        raise bad("Esportazione non disponibile")
    perm, title, sql, cols, ranged = EXPORTS[what]
    ctx.require(perm)
    if what in ("items",) and not ctx.can("inventory.costs"):
        cols = [c_ for c_ in cols if c_[0] not in ("avg_cost", "stock_value", "last_purchase_price")]
    if what == "movements" and not ctx.can("inventory.costs"):
        cols = [c_ for c_ in cols if c_[0] not in ("unit_cost", "value")]
    params = []
    if ranged:
        params = [ctx.qdate("from", dt.date(2000, 1, 1)), ctx.qdate("to", today())]
    with db.read() as c:
        rows = c.all(sql, params)
    content = exporting.table_to_xlsx(title, cols, rows) if fmt == "xlsx" else exporting.table_to_csv(cols, rows)
    mt = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if fmt == "xlsx" else "text/csv; charset=utf-8"
    return Response(content, media_type=mt, headers={"Content-Disposition": f'attachment; filename="{what}-{today()}.{fmt}"'})


FULL_EXCLUDE = {"sessions", "password_resets", "login_attempts", "schema_migrations"}


@api("GET", "/export-all.xlsx", "export.run")
def export_all(ctx: Ctx):
    """Esportazione completa di tutti i dati aziendali (un foglio per tabella). Gli allegati sono esclusi (solo elenco)."""
    wb = Workbook()
    wb.remove(wb.active)
    with db.read() as c:
        tables = [r["table_name"] for r in c.all(
            "SELECT table_name FROM information_schema.tables WHERE table_schema='public' AND table_type='BASE TABLE' ORDER BY 1")]
        for t in tables:
            if t in FULL_EXCLUDE:
                continue
            if t == "attachments":
                rows = c.all("SELECT id, owner_type, owner_id, filename, mime, size_bytes, uploaded_at FROM attachments ORDER BY id")
            elif t == "users":
                rows = c.all("SELECT id, email, full_name, role_id, is_active, created_at, last_login_at FROM users ORDER BY id")
            else:
                rows = c.all(f'SELECT * FROM "{t}" ORDER BY 1')
            ws = wb.create_sheet(t[:31])
            if not rows:
                ws.append(["(nessun dato)"])
                continue
            keys = list(rows[0].keys())
            ws.append(keys)
            for r in rows:
                vals = []
                for k in keys:
                    v = r[k]
                    if isinstance(v, Decimal):
                        v = float(v)
                    elif isinstance(v, dt.datetime):
                        v = v.replace(tzinfo=None)
                    elif isinstance(v, (dict, list)):
                        import json
                        v = json.dumps(v, ensure_ascii=False, default=str)
                    elif isinstance(v, (bytes, memoryview)):
                        v = None
                    vals.append(v)
                ws.append(vals)
    with db.tx() as c:
        audit(c, ctx, "export", "all", None, "Esportazione completa dei dati")
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="isa-manager-dati-completi-{today()}.xlsx"'})
