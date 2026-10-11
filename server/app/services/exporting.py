"""Esportazioni: report in Excel/PDF, tabelle in CSV/Excel."""
from __future__ import annotations

import csv
import datetime as dt
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .. import db
from ..common import D, add_months, get_settings, month_end, month_start, q2
from . import reports as R

EURO_FMT = '#,##0.00 "€";[Red]-#,##0.00 "€"'
BRAND = "2F7F78"


def fmt_eur(v) -> str:
    if v is None:
        return "–"
    s = f"{D(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{s} €"


def fmt_date(d) -> str:
    if isinstance(d, str):
        d = dt.date.fromisoformat(d[:10])
    return d.strftime("%d/%m/%Y") if d else ""


def report_data(c: db.Cursor, d1: dt.date, d2: dt.date) -> dict:
    s = get_settings(c)
    pl = R.profit_and_loss(c, d1, d2)
    months = (d2.year - d1.year) * 12 + d2.month - d1.month + 1
    if months >= 2:
        series = [{"label": m["month"].strftime("%m/%Y"), "revenue": m["revenue"], "costs": m["costs"], "result": m["result"]}
                  for m in R.monthly_series(c, d2, min(months, 24))]
    else:
        series = [{"label": r["day"].strftime("%d"), "revenue": q2(D(r["net"])), "costs": None, "result": None}
                  for r in R.daily_series(c, d1, d2)]
    return {
        "company": s.get("company", {}), "from": d1, "to": d2, "pl": pl,
        "liquidity": R.liquidity(c, d1, d2), "checks": R.consistency_checks(c, d1, d2),
        "expenses": R.expenses_by_category(c, d1, d2), "products": R.product_analysis(c, d1, d2),
        "channels": R.channel_analysis(c, d1, d2), "series": series,
        "generated": dt.datetime.now(),
    }


def _period_label(d1, d2):
    if d1.day == 1 and d2 == month_end(d1):
        mesi = ["gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno", "luglio", "agosto", "settembre",
                "ottobre", "novembre", "dicembre"]
        return f"{mesi[d1.month - 1].capitalize()} {d1.year}"
    return f"dal {fmt_date(d1)} al {fmt_date(d2)}"


DISCLAIMER = ("Stima gestionale elaborata da Isaa Rolls Manager: non costituisce bilancio fiscale ufficiale. "
              "I criteri di contabilizzazione vanno verificati con il commercialista.")


# ------------------------------------------------------------------ Excel

def _header(ws, row, values, widths=None):
    for i, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BRAND)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def pl_to_xlsx(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Conto economico"
    pl, liq = data["pl"], data["liquidity"]
    ws["A1"] = f"{data['company'].get('name') or 'Isaa Rolls Manager'} – Report economico {_period_label(data['from'], data['to'])}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = DISCLAIMER
    ws["A2"].font = Font(italic=True, size=9, color="666666")
    _header(ws, 4, ["Voce", "Importo", "% sui ricavi", "Note"], [52, 16, 12, 40])
    r = 5
    for ln in pl["lines"]:
        ws.cell(row=r, column=1, value=("    " * (ln["level"] - 1 if ln["level"] > 1 else 0)) + ln["label"])
        ws.cell(row=r, column=2, value=float(ln["amount"])).number_format = EURO_FMT
        if ln["pct"] is not None:
            ws.cell(row=r, column=3, value=float(ln["pct"]) / 100).number_format = "0.0%"
        ws.cell(row=r, column=4, value=ln.get("note") or "")
        if ln["kind"] in ("subtotal", "result"):
            for col in range(1, 5):
                ws.cell(row=r, column=col).font = Font(bold=True)
                ws.cell(row=r, column=col).border = Border(top=Side(style="thin"))
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="IVA (stima)").font = Font(bold=True); r += 1
    for k, l in (("output", "IVA a debito sulle vendite"), ("input", "IVA a credito su acquisti e spese"),
                 ("balance", "Saldo IVA stimato (positivo = da versare)")):
        ws.cell(row=r, column=1, value=l)
        ws.cell(row=r, column=2, value=float(pl["vat"][k])).number_format = EURO_FMT
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Liquidità (distinta dal risultato economico)").font = Font(bold=True); r += 1
    for k, l in (("inflows", "Entrate monetarie effettive"), ("outflows", "Uscite monetarie effettive"),
                 ("net_cash_flow", "Flusso di cassa netto"), ("receivables", "Crediti da incassare (fine periodo)"),
                 ("payables", "Debiti da pagare (fine periodo)"), ("cash_total", "Saldo cassa e conti (fine periodo)"),
                 ("stock_value", "Valore scorte di magazzino (fine periodo)")):
        ws.cell(row=r, column=1, value=l)
        ws.cell(row=r, column=2, value=float(liq[k])).number_format = EURO_FMT
        r += 1
    for a in liq["accounts"]:
        ws.cell(row=r, column=1, value=f"    {a['name']}")
        ws.cell(row=r, column=2, value=float(a["balance"])).number_format = EURO_FMT
        r += 1
    if pl.get("excluded"):
        r += 1
        ws.cell(row=r, column=1, value="Importi esclusi dal conto economico (per evitare doppi conteggi)").font = Font(bold=True); r += 1
        labels = {"acquisti_magazzino": "Acquisti di magazzino (costo rilevato al consumo)",
                  "commissioni_fatturate": "Fatture commissioni piattaforme (già stimate sulle vendite)",
                  "non_economiche": "Spese non economiche"}
        for k, v in pl["excluded"].items():
            ws.cell(row=r, column=1, value=labels.get(k, k))
            ws.cell(row=r, column=2, value=float(v)).number_format = EURO_FMT
            r += 1

    ws2 = wb.create_sheet("Andamento")
    _header(ws2, 1, ["Periodo", "Ricavi netti", "Costi", "Risultato"], [14, 16, 16, 16])
    for i, srow in enumerate(data["series"], start=2):
        ws2.cell(row=i, column=1, value=srow["label"])
        for j, k in enumerate(("revenue", "costs", "result"), start=2):
            if srow[k] is not None:
                ws2.cell(row=i, column=j, value=float(srow[k])).number_format = EURO_FMT
    if data["series"]:
        ch = BarChart()
        ch.title = "Ricavi e costi"
        ch.y_axis.title = "€"
        n = len(data["series"]) + 1
        cols = 3 if data["series"][0]["costs"] is not None else 2
        ch.add_data(Reference(ws2, min_col=2, max_col=cols, min_row=1, max_row=n), titles_from_data=True)
        ch.set_categories(Reference(ws2, min_col=1, min_row=2, max_row=n))
        ch.height, ch.width = 9, 22
        ws2.add_chart(ch, "F2")

    ws3 = wb.create_sheet("Spese per categoria")
    _header(ws3, 1, ["Categoria", "Gruppo", "N. documenti", "Importo", "Nel conto economico"], [34, 22, 14, 16, 20])
    for i, e in enumerate(data["expenses"], start=2):
        ws3.cell(row=i, column=1, value=e["category"])
        ws3.cell(row=i, column=2, value=R.GROUP_LABELS.get(e["pl_group"], e["pl_group"]))
        ws3.cell(row=i, column=3, value=e["count"])
        ws3.cell(row=i, column=4, value=float(e["amount"])).number_format = EURO_FMT
        ws3.cell(row=i, column=5, value="sì" if e["in_pl"] else "no")

    ws4 = wb.create_sheet("Prodotti")
    _header(ws4, 1, ["Prodotto", "Categoria", "Quantità", "Vendite lorde", "Imponibile", "Costo teorico", "Margine", "Food cost %"],
            [30, 16, 10, 15, 15, 15, 15, 12])
    for i, p in enumerate(data["products"], start=2):
        ws4.cell(row=i, column=1, value=p["name"])
        ws4.cell(row=i, column=2, value=p["category"])
        ws4.cell(row=i, column=3, value=float(p["qty"]))
        for j, k in enumerate(("gross", "net", "cost", "margin"), start=4):
            ws4.cell(row=i, column=j, value=float(p[k] or 0)).number_format = EURO_FMT
        if p["food_cost_pct"] is not None:
            ws4.cell(row=i, column=8, value=float(p["food_cost_pct"]) / 100).number_format = "0.0%"

    ws5 = wb.create_sheet("Canali")
    _header(ws5, 1, ["Canale", "N. vendite", "Lordo", "Imponibile", "Commissioni", "% commissioni", "Contribuzione"],
            [24, 12, 15, 15, 15, 14, 15])
    for i, ch_ in enumerate(data["channels"], start=2):
        ws5.cell(row=i, column=1, value=ch_["name"])
        ws5.cell(row=i, column=2, value=ch_["n"])
        for j, k in enumerate(("gross", "net", "commission"), start=3):
            ws5.cell(row=i, column=j, value=float(ch_[k])).number_format = EURO_FMT
        if ch_["commission_pct"] is not None:
            ws5.cell(row=i, column=6, value=float(ch_["commission_pct"]) / 100).number_format = "0.0%"
        ws5.cell(row=i, column=7, value=float(ch_["contribution"])).number_format = EURO_FMT

    ws6 = wb.create_sheet("Controlli")
    _header(ws6, 1, ["Controllo", "Esito", "Dettaglio"], [48, 10, 90])
    for i, ck in enumerate(data["checks"], start=2):
        ws6.cell(row=i, column=1, value=ck["label"])
        ws6.cell(row=i, column=2, value="OK" if ck["ok"] else ("ATTENZIONE" if ck.get("warning") else "ERRORE"))
        ws6.cell(row=i, column=3, value=ck["detail"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF

def pl_to_pdf(data: dict) -> bytes:
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    brand = colors.HexColor("#" + BRAND)
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontSize=16, textColor=colors.HexColor("#1b2b2a"), spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=11.5, textColor=brand, spaceBefore=10, spaceAfter=4)
    small = ParagraphStyle("small", parent=ss["Normal"], fontSize=8, textColor=colors.HexColor("#6b625d"))
    body = ParagraphStyle("body", parent=ss["Normal"], fontSize=9)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title=f"Report economico {_period_label(data['from'], data['to'])}")
    pl, liq = data["pl"], data["liquidity"]
    el = [
        Paragraph(f"{data['company'].get('name') or 'Isaa Rolls Manager'} — Report economico", h1),
        Paragraph(f"Periodo: <b>{_period_label(data['from'], data['to'])}</b> · generato il "
                  f"{data['generated'].strftime('%d/%m/%Y %H:%M')}", body),
        Paragraph(DISCLAIMER, small),
        Spacer(1, 4),
    ]

    def table(rows, widths, bold_rows=(), right_cols=(1, 2), zebra=True, header=True):
        t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
        st = [("FONTSIZE", (0, 0), (-1, -1), 8.5), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
              ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]
        if header:
            st += [("BACKGROUND", (0, 0), (-1, 0), brand), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                   ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]
        for c_ in right_cols:
            st.append(("ALIGN", (c_, 0), (c_, -1), "RIGHT"))
        for r_ in bold_rows:
            st += [("FONTNAME", (0, r_), (-1, r_), "Helvetica-Bold"), ("LINEABOVE", (0, r_), (-1, r_), 0.6, colors.HexColor("#1b2b2a"))]
        if zebra:
            for i in range(1 if header else 0, len(rows)):
                if i % 2 == 0:
                    st.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#eef6f5")))
        t.setStyle(TableStyle(st))
        return t

    rows = [["Voce", "Importo", "% ricavi"]]
    bold = []
    for ln in pl["lines"]:
        indent = "&nbsp;" * (4 * (ln["level"] - 1)) if ln["level"] > 1 else ""
        rows.append([Paragraph(indent + ln["label"] + (f" <font size=7 color='#6b625d'>({ln['note']})</font>" if ln.get("note") else ""), body),
                     fmt_eur(ln["amount"]), f"{float(ln['pct']):.1f}%".replace(".", ",") if ln["pct"] is not None else ""])
        if ln["kind"] in ("subtotal", "result"):
            bold.append(len(rows) - 1)
    el += [Paragraph("Conto economico gestionale", h2), table(rows, [112 * mm, 38 * mm, 28 * mm], bold)]

    # grafico
    series = data["series"]
    if series:
        d = Drawing(178 * mm, 62 * mm)
        bc = VerticalBarChart()
        bc.x, bc.y, bc.width, bc.height = 14 * mm, 10 * mm, 160 * mm, 46 * mm
        vals = [[float(s["revenue"] or 0) for s in series]]
        if series[0]["costs"] is not None:
            vals.append([float(s["costs"] or 0) for s in series])
        bc.data = vals
        bc.categoryAxis.categoryNames = [s["label"] for s in series]
        bc.categoryAxis.labels.fontSize = 6.5
        bc.valueAxis.labels.fontSize = 7
        bc.valueAxis.valueMin = 0
        bc.bars[0].fillColor = brand
        if len(vals) > 1:
            bc.bars[1].fillColor = colors.HexColor("#c9b8ab")
        bc.barSpacing = 0.5
        bc.groupSpacing = 3
        d.add(bc)
        d.add(String(14 * mm, 58 * mm, "Ricavi netti (colore pieno)" + (" e costi (grigio)" if len(vals) > 1 else " giornalieri"),
                     fontSize=7.5, fillColor=colors.HexColor("#6b625d")))
        el += [Paragraph("Andamento", h2), d]

    liq_rows = [["Liquidità", "Importo"]] + [[l, fmt_eur(liq[k])] for k, l in (
        ("inflows", "Entrate monetarie effettive"), ("outflows", "Uscite monetarie effettive"),
        ("net_cash_flow", "Flusso di cassa netto"), ("receivables", "Crediti da incassare (fine periodo)"),
        ("payables", "Debiti da pagare (fine periodo)"), ("cash_total", "Saldo cassa e conti (fine periodo)"),
        ("stock_value", "Valore scorte di magazzino (fine periodo)"))]
    liq_rows += [[f"   {a['name']}", fmt_eur(a["balance"])] for a in liq["accounts"]]
    vat_rows = [["IVA (stima)", "Importo"], ["IVA a debito sulle vendite", fmt_eur(pl["vat"]["output"])],
                ["IVA a credito su acquisti e spese", fmt_eur(pl["vat"]["input"])],
                ["Saldo stimato (positivo = da versare)", fmt_eur(pl["vat"]["balance"])]]
    el += [KeepTogether([Paragraph("Liquidità e posizione finanziaria", h2), table(liq_rows, [140 * mm, 38 * mm], right_cols=(1,))]),
           KeepTogether([Paragraph("IVA", h2), table(vat_rows, [140 * mm, 38 * mm], [3], right_cols=(1,))])]

    if data["expenses"]:
        er = [["Categoria di spesa", "Importo", "Nel C.E."]] + [
            [e["category"], fmt_eur(e["amount"]), "sì" if e["in_pl"] else "no"] for e in data["expenses"]]
        el += [KeepTogether([Paragraph("Spese per categoria", h2), table(er, [118 * mm, 38 * mm, 22 * mm])])]
    if data["products"]:
        pr = [["Prodotto", "Qtà", "Imponibile", "Costo teorico", "Margine", "Food cost"]] + [
            [p["name"], f"{float(p['qty']):g}".replace(".", ","), fmt_eur(p["net"]), fmt_eur(p["cost"]), fmt_eur(p["margin"]),
             f"{float(p['food_cost_pct']):.1f}%".replace(".", ",") if p["food_cost_pct"] is not None else "–"]
            for p in data["products"][:25]]
        el += [Paragraph("Prodotti venduti", h2), table(pr, [58 * mm, 16 * mm, 28 * mm, 28 * mm, 28 * mm, 20 * mm],
                                                        right_cols=(1, 2, 3, 4, 5))]
    ck = [["Controllo di coerenza", "Esito"]] + [
        [Paragraph(f"{c_['label']}<br/><font size=7 color='#6b625d'>{c_['detail']}</font>", body),
         "OK" if c_["ok"] else ("Attenzione" if c_.get("warning") else "ERRORE")] for c_ in data["checks"]]
    el += [KeepTogether([Paragraph("Controlli automatici", h2), table(ck, [150 * mm, 28 * mm], right_cols=())])]
    crit = pl["criteria"]
    el += [Spacer(1, 6), Paragraph(
        f"Criteri applicati: costo materie prime calcolato sui <b>{crit['materials']}</b>; commissioni delle piattaforme da "
        f"<b>{crit['commissions']}</b>; regime IVA <b>{crit['vat_regime']}</b>; valorizzazione scorte a costo medio ponderato."
        + (" Costi fissi e del personale ripartiti per giorni perché il periodo non copre mesi interi." if pl["prorated"] else ""), small)]

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#6b625d"))
        canvas.drawString(16 * mm, 8 * mm, "Isaa Rolls Manager · stima gestionale")
        canvas.drawRightString(A4[0] - 16 * mm, 8 * mm, f"Pagina {doc_.page}")
        canvas.restoreState()

    doc.build(el, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ------------------------------------------------------------------ tabelle generiche

def table_to_xlsx(title: str, columns: list[tuple[str, str]], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    _header(ws, 1, [l for _, l in columns], [max(12, min(40, len(l) + 6)) for _, l in columns])
    for i, r in enumerate(rows, start=2):
        for j, (k, _) in enumerate(columns, start=1):
            v = r.get(k)
            if isinstance(v, Decimal):
                v = float(v)
            elif isinstance(v, dt.datetime):
                v = v.replace(tzinfo=None)
            elif isinstance(v, (dict, list)):
                v = str(v)
            ws.cell(row=i, column=j, value=v)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def table_to_csv(columns: list[tuple[str, str]], rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([l for _, l in columns])
    for r in rows:
        out = []
        for k, _ in columns:
            v = r.get(k)
            if isinstance(v, Decimal):
                v = str(v).replace(".", ",")
            elif isinstance(v, float):
                v = str(v).replace(".", ",")
            elif isinstance(v, dt.datetime):
                v = v.strftime("%d/%m/%Y %H:%M")
            elif isinstance(v, dt.date):
                v = v.strftime("%d/%m/%Y")
            elif v is None:
                v = ""
            out.append(v)
        w.writerow(out)
    return ("﻿" + buf.getvalue()).encode("utf-8")
